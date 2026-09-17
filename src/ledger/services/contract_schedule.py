"""Propose, confirm, and chart contract schedule rows (D47, D48)."""

from __future__ import annotations

import calendar
import re
from datetime import UTC, date, datetime, timedelta

from docx import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import FileNotDecryptedError, PdfReadError
from sqlalchemy import select
from sqlalchemy.orm import Session

from ledger.models import Award, Document
from ledger.models.contract_schedule import ScheduleItem, ScheduleKind
from ledger.models.documents import ComplianceStatus
from ledger.schemas.contract_schedule import (
    GanttBarOut,
    GanttOut,
    ScheduleConfirmIn,
    ScheduleDraftIn,
    ScheduleItemOut,
    SchedulePatch,
    ScheduleProposeIn,
    ScheduleProposeOut,
)
from ledger.services.audit import record_event
from ledger.services.documents import stored_path

EXTRACTABLE_EXT = frozenset({".txt", ".csv", ".docx", ".pdf"})
_KEYWORD = re.compile(
    r"deliverable|milestone|due|report|cdrl|\bsow\b|statement of work",
    re.IGNORECASE,
)
_ISO = re.compile(r"\b((?:19|20)\d{2}-\d{2}-\d{2})\b")
_US = re.compile(r"\b([A-Za-z]+)\s+(\d{1,2}),\s*((?:19|20)\d{2})\b")
_NUMERIC = re.compile(r"\b(\d{1,2})[/-](\d{1,2})[/-]((?:19|20)\d{2})\b")
_DAY_MONTH = re.compile(r"\b(\d{1,2})[\s-]+([A-Za-z]+)[\s,-]+((?:19|20)\d{2})\b")
_MONTH_YEAR = re.compile(r"\b([A-Za-z]+)\s+((?:19|20)\d{2})\b")
_RELATIVE_MONTH = re.compile(
    r"\b(?:month\s*(\d+)|(\d+)\s*months?\s*(?:after|from)\s*"
    r"(?:award|award start|start|kickoff))\b",
    re.IGNORECASE,
)
_MONTHS = {
    **{calendar.month_name[i].lower(): i for i in range(1, 13)},
    **{calendar.month_abbr[i].lower(): i for i in range(1, 13)},
}


class ScheduleTrackError(ValueError):
    """Domain error for contract schedule."""


def _iso_date(value: str, *, field: str) -> str:
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError as exc:
        raise ScheduleTrackError(f"{field} must be YYYY-MM-DD") from exc


def _parse_date(value: str) -> date:
    return date.fromisoformat(_iso_date(value, field="date"))


def _today() -> date:
    return datetime.now(UTC).date()


def require_award(session: Session, award_id: int) -> Award:
    """Load an award or raise."""
    award = session.get(Award, award_id)
    if award is None:
        raise ScheduleTrackError("award not found")
    return award


def require_kind(session: Session, kind_code: str) -> str:
    """Validate schedule_kind."""
    row = session.get(ScheduleKind, kind_code)
    if row is None:
        raise ScheduleTrackError("unknown schedule kind")
    return kind_code


def require_status(session: Session, status_code: str) -> str:
    """Reuse compliance_status."""
    row = session.get(ComplianceStatus, status_code)
    if row is None:
        raise ScheduleTrackError("unknown status")
    return status_code


def serialize_item(row: ScheduleItem) -> ScheduleItemOut:
    """API shape for a stored row."""
    return ScheduleItemOut(
        schedule_item_id=row.schedule_item_id,
        award_id=row.award_id,
        kind_code=row.kind_code,
        title=row.title,
        start_date=row.start_date,
        due_date=row.due_date,
        status_code=row.status_code,
        notes=row.notes,
        completed_at=row.completed_at,
        source_document_id=row.source_document_id,
        origin_code=row.origin_code,
        created_at=row.created_at,
    )


def list_schedule(session: Session, award_id: int) -> list[ScheduleItem]:
    """Items on one award, due date then id."""
    return list(
        session.scalars(
            select(ScheduleItem)
            .where(ScheduleItem.award_id == award_id)
            .order_by(ScheduleItem.due_date, ScheduleItem.schedule_item_id)
        )
    )


def _lerp(start: date, end: date, numerator: int, denominator: int) -> date:
    span = (end - start).days
    return start + timedelta(days=span * numerator // denominator)


def _draft(
    award: Award,
    *,
    kind_code: str,
    title: str,
    start: date | None,
    due: date,
    origin_code: str,
    notes: str | None = None,
    source_document_id: int | None = None,
) -> ScheduleItemOut:
    return ScheduleItemOut(
        award_id=award.award_id,
        kind_code=kind_code,
        title=title,
        start_date=start.isoformat() if start else None,
        due_date=due.isoformat(),
        status_code="open",
        notes=notes,
        completed_at=None,
        source_document_id=source_document_id,
        origin_code=origin_code,
    )


def template_drafts(award: Award) -> list[ScheduleItemOut]:
    """Starter rows from PoP and phase (D47)."""
    start = _parse_date(award.pop_start)
    end = _parse_date(award.pop_end)
    items = [
        _draft(
            award,
            kind_code="milestone",
            title="Kickoff",
            start=start,
            due=start,
            origin_code="template",
        ),
        _draft(
            award,
            kind_code="pop",
            title="Period of performance",
            start=start,
            due=end,
            origin_code="template",
        ),
    ]
    phase = award.phase_code
    if phase == "I":
        items.append(
            _draft(
                award,
                kind_code="report",
                title="Interim technical report",
                start=start,
                due=_lerp(start, end, 1, 2),
                origin_code="template",
            )
        )
        items.append(
            _draft(
                award,
                kind_code="deliverable",
                title="Final technical report / deliverable",
                start=start,
                due=end,
                origin_code="template",
            )
        )
    elif phase in {"II", "IIB"}:
        for num, label in (
            (1, "Q1 progress report"),
            (2, "Mid-period review"),
            (3, "Q3 progress report"),
        ):
            items.append(
                _draft(
                    award,
                    kind_code="report",
                    title=label,
                    start=start,
                    due=_lerp(start, end, num, 4),
                    origin_code="template",
                )
            )
        items.append(
            _draft(
                award,
                kind_code="deliverable",
                title="Final report / deliverable",
                start=start,
                due=end,
                origin_code="template",
            )
        )
    else:
        items.append(
            _draft(
                award,
                kind_code="report",
                title="Mid-period report",
                start=start,
                due=_lerp(start, end, 1, 2),
                origin_code="template",
            )
        )
        items.append(
            _draft(
                award,
                kind_code="deliverable",
                title="Final deliverable",
                start=start,
                due=end,
                origin_code="template",
            )
        )
    return items


def _us_date(month_name: str, day: str, year: str) -> date | None:
    month = _MONTHS.get(month_name.lower())
    if month is None:
        return None
    try:
        return date(int(year), month, int(day))
    except ValueError:
        return None


def _add_months(value: date, months: int) -> date:
    """Add calendar months, clamping the day to the target month."""
    month_index = value.year * 12 + value.month - 1 + months
    year, zero_based_month = divmod(month_index, 12)
    month = zero_based_month + 1
    day = min(value.day, calendar.monthrange(year, month)[1])
    return date(year, month, day)


def _dates_in_line(line: str, award: Award) -> list[date]:
    """Recognize common absolute and award-relative contract dates."""
    dates: list[date] = []
    spans: list[tuple[int, int]] = []

    def add(parsed: date | None, span: tuple[int, int]) -> None:
        overlaps = any(span[0] < existing[1] and span[1] > existing[0] for existing in spans)
        if parsed is not None and not overlaps:
            dates.append(parsed)
            spans.append(span)

    for match in _ISO.finditer(line):
        try:
            add(date.fromisoformat(match.group(1)), match.span())
        except ValueError:
            continue
    for match in _US.finditer(line):
        add(_us_date(match.group(1), match.group(2), match.group(3)), match.span())
    for match in _NUMERIC.finditer(line):
        try:
            add(date(int(match.group(3)), int(match.group(1)), int(match.group(2))), match.span())
        except ValueError:
            continue
    for match in _DAY_MONTH.finditer(line):
        add(_us_date(match.group(2), match.group(1), match.group(3)), match.span())
    for match in _MONTH_YEAR.finditer(line):
        month = _MONTHS.get(match.group(1).lower())
        if month is not None:
            year = int(match.group(2))
            add(date(year, month, calendar.monthrange(year, month)[1]), match.span())
    award_start = _parse_date(award.pop_start)
    for match in _RELATIVE_MONTH.finditer(line):
        count = int(match.group(1) or match.group(2))
        add(_add_months(award_start, count), match.span())
    return dates


def extract_dated_lines(
    text: str,
    award: Award,
    *,
    source_document_id: int | None,
) -> list[ScheduleItemOut]:
    """Heuristic dated deliverable/milestone lines from contract text."""
    found: list[ScheduleItemOut] = []
    seen: set[tuple[str, str]] = set()
    lines = [" ".join(raw.split()) for raw in text.splitlines()]
    lines = [line for line in lines if line]
    for index, line in enumerate(lines):
        if not line or not _KEYWORD.search(line):
            continue
        dates = _dates_in_line(line, award)
        if not dates:
            # PDF extraction and Word tables often put the title and date on
            # adjacent lines/cells. Keep the context narrow to avoid joining
            # unrelated schedule clauses.
            for following in lines[index + 1 : index + 3]:
                line = f"{line} {following}"
                dates = _dates_in_line(line, award)
                if dates:
                    break
                if _KEYWORD.search(following):
                    break
        if not dates:
            continue
        due = dates[-1]
        title = line
        if len(title) > 160:
            title = title[:157] + "..."
        kind = "report" if re.search(r"report", line, re.IGNORECASE) else "deliverable"
        if re.search(r"milestone", line, re.IGNORECASE):
            kind = "milestone"
        key = (title.lower(), due.isoformat())
        if key in seen:
            continue
        seen.add(key)
        found.append(
            _draft(
                award,
                kind_code=kind,
                title=title,
                start=_parse_date(award.pop_start),
                due=due,
                origin_code="extract",
                notes="Extracted from contract text; confirm before saving.",
                source_document_id=source_document_id,
            )
        )
    return found


def _read_document_text(row: Document) -> tuple[str | None, str | None]:
    """Return (text, skip_reason)."""
    path = stored_path(row)
    if path is None or not path.is_file():
        return None, "document has no file on disk"
    ext = (row.stored_ext or path.suffix or "").lower()
    if not ext.startswith("."):
        ext = f".{ext}" if ext else ""
    if ext == ".doc":
        return None, "Legacy .doc files cannot be parsed; convert to .docx or paste SOW text"
    if ext not in EXTRACTABLE_EXT:
        return None, "This file type cannot be parsed; paste SOW text or confirm the template"
    try:
        if ext in {".txt", ".csv"}:
            text = path.read_text(encoding="utf-8", errors="replace")
        elif ext == ".docx":
            document = WordDocument(path)
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row_cells in table.rows:
                    line = " | ".join(
                        cell.text.strip() for cell in row_cells.cells if cell.text.strip()
                    )
                    if line:
                        parts.append(line)
            text = "\n".join(parts)
        else:
            reader = PdfReader(path)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except (OSError, PackageNotFoundError, PdfReadError, FileNotDecryptedError):
        return None, "could not read document file"
    if not text.strip():
        if ext == ".pdf":
            return None, "PDF has no extractable text; it may be scanned. Paste SOW text instead"
        return None, "document contains no extractable text"
    return text, None


def propose_schedule(
    session: Session,
    award: Award,
    payload: ScheduleProposeIn,
) -> ScheduleProposeOut:
    """Build a draft. Does not insert rows."""
    notes: list[str] = []
    items = template_drafts(award)
    source_id = payload.document_id
    if source_id is not None:
        doc = session.get(Document, source_id)
        if doc is None or doc.award_id != award.award_id:
            raise ScheduleTrackError("document not found on this award")
        text, skip = _read_document_text(doc)
        if skip:
            notes.append(skip)
        elif text:
            extracted = extract_dated_lines(text, award, source_document_id=source_id)
            if extracted:
                items.extend(extracted)
            else:
                notes.append("no dated deliverable lines found in the file")
    pasted = (payload.text or "").strip()
    if pasted:
        extracted = extract_dated_lines(pasted, award, source_document_id=source_id)
        if extracted:
            items.extend(extracted)
        else:
            notes.append("no dated deliverable lines found in the pasted text")
    notes.append("Nothing is saved until you confirm the rows you want to keep.")
    return ScheduleProposeOut(award_id=award.award_id, notes=notes, items=items)


def _insert_item(
    session: Session,
    award: Award,
    draft: ScheduleDraftIn,
    *,
    actor_id: int | None,
) -> ScheduleItem:
    kind = require_kind(session, draft.kind_code)
    origin = (
        draft.origin_code if draft.origin_code in {"template", "extract", "manual"} else "manual"
    )
    start = _iso_date(draft.start_date, field="start_date") if draft.start_date else None
    due = _iso_date(draft.due_date, field="due_date")
    if start and due < start:
        raise ScheduleTrackError("due_date must be on or after start_date")
    source_id = draft.source_document_id
    if source_id is not None:
        doc = session.get(Document, source_id)
        if doc is None or doc.award_id != award.award_id:
            raise ScheduleTrackError("document not found on this award")
    row = ScheduleItem(
        award_id=award.award_id,
        kind_code=kind,
        title=draft.title.strip(),
        start_date=start,
        due_date=due,
        status_code="open",
        notes=draft.notes,
        source_document_id=source_id,
        origin_code=origin,
        created_by=actor_id,
    )
    if not row.title:
        raise ScheduleTrackError("title is required")
    session.add(row)
    session.flush()
    record_event(
        session,
        action="schedule_create",
        entity_type="schedule_item",
        entity_id=row.schedule_item_id,
        actor_user_id=actor_id,
        detail={"award_id": award.award_id, "title": row.title},
    )
    return row


def confirm_schedule(
    session: Session,
    award: Award,
    payload: ScheduleConfirmIn,
    *,
    actor_id: int | None,
) -> list[ScheduleItem]:
    """Insert kept draft rows."""
    created: list[ScheduleItem] = []
    for draft in payload.items:
        if not draft.keep:
            continue
        created.append(_insert_item(session, award, draft, actor_id=actor_id))
    if created:
        record_event(
            session,
            action="schedule_confirm",
            entity_type="award",
            entity_id=award.award_id,
            actor_user_id=actor_id,
            detail={"count": len(created)},
        )
    return created


def create_schedule_item(
    session: Session,
    award: Award,
    draft: ScheduleDraftIn,
    *,
    actor_id: int | None,
) -> ScheduleItem:
    """Manual add."""
    draft.origin_code = "manual"
    return _insert_item(session, award, draft, actor_id=actor_id)


def patch_schedule_item(
    session: Session,
    row: ScheduleItem,
    payload: SchedulePatch,
    *,
    actor_id: int | None,
) -> ScheduleItem:
    """Edit title/dates/status."""
    if payload.title is not None:
        title = payload.title.strip()
        if not title:
            raise ScheduleTrackError("title is required")
        row.title = title
    if payload.kind_code is not None:
        row.kind_code = require_kind(session, payload.kind_code)
    if payload.start_date is not None:
        row.start_date = (
            _iso_date(payload.start_date, field="start_date") if payload.start_date else None
        )
    if payload.due_date is not None:
        row.due_date = _iso_date(payload.due_date, field="due_date")
    start = row.start_date
    if start and row.due_date < start:
        raise ScheduleTrackError("due_date must be on or after start_date")
    if payload.notes is not None:
        row.notes = payload.notes
    if payload.status_code is not None:
        status_code = require_status(session, payload.status_code)
        row.status_code = status_code
        if status_code in {"done", "waived"} and not row.completed_at:
            row.completed_at = datetime.now(UTC).replace(microsecond=0).isoformat()
        if status_code == "open":
            row.completed_at = None
        record_event(
            session,
            action="schedule_status",
            entity_type="schedule_item",
            entity_id=row.schedule_item_id,
            actor_user_id=actor_id,
            detail={"status_code": status_code},
        )
    else:
        record_event(
            session,
            action="schedule_update",
            entity_type="schedule_item",
            entity_id=row.schedule_item_id,
            actor_user_id=actor_id,
        )
    session.flush()
    return row


def delete_schedule_item(session: Session, row: ScheduleItem, *, actor_id: int | None) -> None:
    """Remove a schedule row (not remaining money)."""
    item_id = row.schedule_item_id
    session.delete(row)
    session.flush()
    record_event(
        session,
        action="schedule_delete",
        entity_type="schedule_item",
        entity_id=item_id,
        actor_user_id=actor_id,
    )


def _lane(row: ScheduleItem, as_of: date) -> str:
    if row.status_code in {"done", "waived"}:
        return "completed"
    due = date.fromisoformat(row.due_date)
    if due < as_of:
        return "behind"
    return "remaining"


def _pct(part: int, whole: int) -> int:
    if whole <= 0:
        return 0
    return max(0, min(100, (part * 100) // whole))


def build_gantt(
    session: Session,
    *,
    award_id: int | None,
    as_of: date | None = None,
) -> GanttOut:
    """Compute bars from confirmed schedule rows."""
    as_of = as_of or _today()
    query = select(ScheduleItem).order_by(ScheduleItem.due_date, ScheduleItem.schedule_item_id)
    if award_id is not None:
        require_award(session, award_id)
        query = query.where(ScheduleItem.award_id == award_id)
    rows = list(session.scalars(query))
    if not rows:
        return GanttOut(as_of=as_of.isoformat(), chart_start=None, chart_end=None, bars=[])
    awards: dict[int, Award] = {}
    starts: list[date] = []
    ends: list[date] = []
    for row in rows:
        award = awards.get(row.award_id)
        if award is None:
            loaded = session.get(Award, row.award_id)
            if loaded is None:
                continue
            award = loaded
            awards[row.award_id] = award
        bar_start = (
            date.fromisoformat(row.start_date)
            if row.start_date
            else date.fromisoformat(award.pop_start)
        )
        bar_end = date.fromisoformat(row.due_date)
        bar_end = max(bar_end, bar_start)
        starts.append(bar_start)
        ends.append(bar_end)
        starts.append(date.fromisoformat(award.pop_start))
        ends.append(date.fromisoformat(award.pop_end))
    chart_start = min(starts)
    chart_end = max(ends)
    whole = (chart_end - chart_start).days
    bars: list[GanttBarOut] = []
    for row in rows:
        award = awards.get(row.award_id)
        if award is None:
            continue
        bar_start = (
            date.fromisoformat(row.start_date)
            if row.start_date
            else date.fromisoformat(award.pop_start)
        )
        bar_end = date.fromisoformat(row.due_date)
        bar_end = max(bar_end, bar_start)
        offset = (bar_start - chart_start).days
        width = max(1, (bar_end - bar_start).days)
        bars.append(
            GanttBarOut(
                schedule_item_id=row.schedule_item_id,
                award_id=row.award_id,
                award_short_code=award.short_code,
                title=row.title,
                kind_code=row.kind_code,
                start_date=bar_start.isoformat(),
                due_date=bar_end.isoformat(),
                status_code=row.status_code,
                lane=_lane(row, as_of),
                offset_pct=_pct(offset, whole),
                width_pct=max(1, _pct(width, whole)),
            )
        )
    return GanttOut(
        as_of=as_of.isoformat(),
        chart_start=chart_start.isoformat(),
        chart_end=chart_end.isoformat(),
        bars=bars,
    )
