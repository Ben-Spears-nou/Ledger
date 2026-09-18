"""Phase 12: glossary, contract schedule propose/confirm, Gantt (D46–D48)."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from types import SimpleNamespace

from docx import Document as WordDocument
from fastapi.testclient import TestClient
from tests.conftest import auth_header, login
from tests.test_phase2_time import _award, _employee
from tests.test_phase5_documents import D18_KEYS, _remaining_approved

from ledger.config import PROJECT_ROOT
from ledger.services import contract_schedule


def test_alembic_head_is_phase12() -> None:
    versions = {path.name for path in (PROJECT_ROOT / "alembic" / "versions").glob("*.py")}
    assert "0009_phase12_schedule_gantt.py" in versions


def test_glossary_and_search_common_language(client: TestClient) -> None:
    admin = login(client)
    employee, _pid = _employee(client, admin, username="pat")
    terms = client.get("/glossary", headers=auth_header(employee))
    assert terms.status_code == 200, terms.text
    codes = {row["term_code"] for row in terms.json()}
    assert {"remaining", "schedule", "gantt", "award"} <= codes

    found = client.get("/search", params={"q": "whats left"}, headers=auth_header(admin))
    assert found.status_code == 200, found.text
    kinds = {hit["kind"] for hit in found.json()["hits"]}
    assert "glossary" in kinds
    remaining = [hit for hit in found.json()["hits"] if hit["kind"] == "glossary"]
    assert any("Remaining" in hit["label"] for hit in remaining)
    assert remaining[0]["href"].startswith("/help")

    denied = client.get("/search", params={"q": "remaining"}, headers=auth_header(employee))
    assert denied.status_code == 403


def test_propose_does_not_write_confirm_does_not_move_remaining(client: TestClient) -> None:
    admin = login(client)
    award = _award(
        client, admin, short_code="P12A", type_code="CPFF", template="CPFF_SBIR", oh_pct=3000
    )
    award_id = award["award_id"]
    before = _remaining_approved(client, admin, award_id)
    listed = client.get(f"/awards/{award_id}/schedule", headers=auth_header(admin))
    assert listed.status_code == 200
    assert listed.json() == []

    proposed = client.post(
        f"/awards/{award_id}/schedule/propose",
        json={},
        headers=auth_header(admin),
    )
    assert proposed.status_code == 200, proposed.text
    items = proposed.json()["items"]
    assert items
    assert all(row["origin_code"] == "template" for row in items)
    still = client.get(f"/awards/{award_id}/schedule", headers=auth_header(admin))
    assert still.json() == []

    confirmed = client.post(
        f"/awards/{award_id}/schedule/confirm",
        json={"items": [{**row, "keep": True} for row in items]},
        headers=auth_header(admin),
    )
    assert confirmed.status_code == 201, confirmed.text
    assert len(confirmed.json()) == len(items)
    assert _remaining_approved(client, admin, award_id) == before

    employee, _pid = _employee(client, admin, username="kim")
    slim = client.get("/awards", headers=auth_header(employee))
    assert slim.status_code == 200
    slim_card = next(row for row in slim.json() if row["award_id"] == award_id)
    assert set(slim_card) == D18_KEYS


def test_extract_from_txt_contract_and_gantt_lanes(client: TestClient) -> None:
    admin = login(client)
    employee, _pid = _employee(client, admin, username="lee")
    award = _award(
        client, admin, short_code="P12B", type_code="FFP", template="FFP_INTERNAL", oh_pct=0
    )
    award_id = award["award_id"]
    created = client.post(
        f"/awards/{award_id}/documents",
        json={"kind_code": "contract", "title": "SOW"},
        headers=auth_header(admin),
    )
    doc_id = created.json()["document_id"]
    upload = client.post(
        f"/documents/{doc_id}/file",
        files={
            "file": (
                "sow.txt",
                b"Deliverable 1: Prototype due 2026-03-15\nIgnore this line\n",
                "text/plain",
            )
        },
        headers=auth_header(admin),
    )
    assert upload.status_code == 200, upload.text

    proposed = client.post(
        f"/awards/{award_id}/schedule/propose",
        json={"document_id": doc_id},
        headers=auth_header(admin),
    )
    assert proposed.status_code == 200, proposed.text
    extracted = [row for row in proposed.json()["items"] if row["origin_code"] == "extract"]
    assert extracted
    assert any("2026-03-15" == row["due_date"] for row in extracted)

    kept = [row for row in extracted if row["due_date"] == "2026-03-15"][:1]
    kept[0]["keep"] = True
    saved = client.post(
        f"/awards/{award_id}/schedule/confirm",
        json={"items": kept},
        headers=auth_header(admin),
    )
    assert saved.status_code == 201, saved.text
    item_id = saved.json()[0]["schedule_item_id"]

    behind = client.get(
        f"/awards/{award_id}/gantt",
        params={"as_of": "2026-03-16"},
        headers=auth_header(admin),
    )
    assert behind.status_code == 200, behind.text
    assert behind.json()["bars"][0]["lane"] == "behind"

    done = client.patch(
        f"/schedule/{item_id}",
        json={"status_code": "done"},
        headers=auth_header(admin),
    )
    assert done.status_code == 200
    completed = client.get(
        "/gantt",
        params={"as_of": "2026-03-16"},
        headers=auth_header(admin),
    )
    assert completed.status_code == 200
    bar = next(row for row in completed.json()["bars"] if row["schedule_item_id"] == item_id)
    assert bar["lane"] == "completed"

    viewed = client.get("/gantt", headers=auth_header(employee))
    assert viewed.status_code == 200, viewed.text
    assert any(row["schedule_item_id"] == item_id for row in viewed.json()["bars"])
    assert (
        client.post(
            f"/awards/{award_id}/schedule/propose",
            json={},
            headers=auth_header(employee),
        ).status_code
        == 403
    )


def test_extract_common_contract_dates_and_wrapped_lines() -> None:
    award = SimpleNamespace(award_id=7, pop_start="2026-01-01")
    text = """\
Deliverable A due 6/15/2026
Milestone B due 15 July 2026
Progress report due August 2026
CDRL A004 due Month 6 after award
Deliverable E
September 15, 2026
"""
    rows = contract_schedule.extract_dated_lines(text, award, source_document_id=12)

    assert [row.due_date for row in rows] == [
        "2026-06-15",
        "2026-07-15",
        "2026-08-31",
        "2026-07-01",
        "2026-09-15",
    ]
    assert all(row.origin_code == "extract" for row in rows)
    assert all(row.source_document_id == 12 for row in rows)
    assert [row.start_date for row in rows] == [
        "2026-06-15",
        "2026-07-15",
        "2026-08-31",
        "2026-07-01",
        "2026-09-15",
    ]


def test_extract_contract_writing_system_clin_schedule() -> None:
    award = SimpleNamespace(
        award_id=7,
        pop_start="2026-09-30",
        pop_end="2027-09-30",
    )
    text = """\
Section B - Supplies or Services & Prices or Costs
0001 Monthly Technical Report -- The US Army Research Office
1 Each USD 58,208.81
0002 Monthly Technical Report -- The US Army Research Office
1 Each USD 58,208.76
Section F - Deliveries or Performance
0001 1 Each Quantity
Address and POC
Period of Performance
From
30 Sep 2026
W911NF26CA035
Page 38 of 100
To
30 Oct 2026
0002 1 Each Quantity
Address and POC
Period of Performance
From
31 Oct 2026
To
30 Nov 2026
"""

    assert contract_schedule.contract_pop_dates(text) == (
        date(2026, 9, 30),
        date(2026, 11, 30),
    )
    rows = contract_schedule.extract_clin_schedule(text, award, source_document_id=12)

    assert [row.title for row in rows] == [
        "CLIN 0001 — Monthly Technical Report",
        "CLIN 0002 — Monthly Technical Report",
    ]
    assert [row.start_date for row in rows] == ["2026-09-30", "2026-10-31"]
    assert [row.due_date for row in rows] == ["2026-10-30", "2026-11-30"]
    assert all(row.kind_code == "report" for row in rows)
    assert all(row.source_document_id == 12 for row in rows)


def test_extract_relative_month_ranges_and_schedule_context() -> None:
    award = SimpleNamespace(
        award_id=7,
        pop_start="2026-09-30",
        pop_end="2028-09-29",
    )
    tasks = """\
Schedule of planned tasks
4.2.2: Procure Materials: Months 1-3
4.2.14: Assess Shelf-Life: Months 4–24
"""
    task_rows = contract_schedule.extract_dated_lines(tasks, award, source_document_id=None)

    assert [row.start_date for row in task_rows] == ["2026-09-30", "2026-12-30"]
    assert [row.due_date for row in task_rows] == ["2026-12-29", "2028-09-29"]

    vertical_pop = """\
Period of Performance
From
30 Sep 2026
To
30 Oct 2026
"""
    pop_rows = contract_schedule.extract_dated_lines(
        vertical_pop,
        award,
        source_document_id=None,
    )
    assert [row.due_date for row in pop_rows] == ["2026-09-30", "2026-10-30"]


def test_extract_cdrl_dac_eoc_and_recurring_schedule() -> None:
    award = SimpleNamespace(
        award_id=7,
        pop_start="2026-01-01",
        pop_end="2026-12-31",
    )
    text = """\
Section F - Deliveries or Performance
0001 Period of Performance
From
31 Aug 2026
To
30 Aug 2028
1. DATA ITEM NO.
A001
2. TITLE OF DATA ITEM
Program Management Plan
10. FREQUENCY
One time
12. DATE OF FIRST SUBMISSION
45 DAC
16. REMARKS
Blk 13: Updated quarterly after initial baseline.
15.TOTAL
1. DATA ITEM NO.
A002
2. TITLE OF DATA ITEM
Progress Report
10. FREQUENCY
Monthly
12. DATE OF FIRST SUBMISSION
30 DAC
16. REMARKS
BLK 13: Monthly Reports shall be submitted NLT 15 days after the end of each month.
15.TOTAL
1. DATA ITEM NO.
A003
2. TITLE OF DATA ITEM
Six Month Project Reviews
10. FREQUENCY
Every 6 months
12. DATE OF FIRST SUBMISSION
180 DAC
16. REMARKS
DAC: Days After Contract Award.
15.TOTAL
1. DATA ITEM NO.
A005
2. TITLE OF DATA ITEM
Technical Data Package
10. FREQUENCY
ASREQ
12. DATE OF FIRST SUBMISSION
EOC
16. REMARKS
EOC: End of Contract
15.TOTAL
1. DATA ITEM NO.
A007
2. TITLE OF DATA ITEM
Final Report
10. FREQUENCY
See BLK 16
12. DATE OF FIRST SUBMISSION
See BLK 16
16. REMARKS
The contractor shall deliver a draft NLT thirty (30) days before the end of the POP.
The contractor shall deliver the final report NLT the end of the POP.
15.TOTAL
1. DATA ITEM NO.
A008
2. TITLE OF DATA ITEM
Patents - Reporting of Subject Inventions
10. FREQUENCY
See BLK 16
12. DATE OF FIRST SUBMISSION
See BLK 16
16. REMARKS
Submit DD Form 882 every 12 months from the date of the contract award.
Submit DD Form 882 in a final report during the contract term.
15.TOTAL
1. DATA ITEM NO.
A009
2. TITLE OF DATA ITEM
One-Time Demonstration
10. FREQUENCY
One time
12. DATE OF FIRST SUBMISSION
15 Oct 2026
16. REMARKS
Point deliverable.
15.TOTAL
"""
    assert contract_schedule.contract_pop_dates(text) == (
        date(2026, 8, 31),
        date(2028, 8, 30),
    )

    rows = contract_schedule.extract_cdrl_schedule(text, award, source_document_id=12)
    by_title = {row.title: row for row in rows}

    assert by_title["A001 Program Management Plan — Baseline"].due_date == "2026-10-15"
    assert by_title["A001 Program Management Plan — Baseline"].start_date == "2026-08-31"
    assert by_title["A001 Program Management Plan — Quarterly update 1"].due_date == "2027-01-15"
    assert by_title["A001 Program Management Plan — Quarterly update 1"].start_date == "2026-10-15"
    assert by_title["A002 Progress Report — Submission 1"].due_date == "2026-09-30"
    assert by_title["A002 Progress Report — Submission 1"].start_date == "2026-08-31"
    assert by_title["A002 Progress Report — Submission 2"].due_date == "2026-10-15"
    assert by_title["A002 Progress Report — Submission 2"].start_date == "2026-09-30"
    assert by_title["A002 Progress Report — Submission 24"].due_date == "2028-08-15"
    assert by_title["A003 Six Month Project Reviews — Review 1"].due_date == "2027-02-27"
    assert by_title["A003 Six Month Project Reviews — Review 1"].start_date == "2026-08-31"
    assert by_title["A005 Technical Data Package"].due_date == "2028-08-30"
    assert by_title["A005 Technical Data Package"].start_date == "2028-08-30"
    assert by_title["A007 Final Report — Draft"].due_date == "2028-07-31"
    assert by_title["A007 Final Report — Draft"].start_date == "2028-07-31"
    assert by_title["A007 Final Report — Final"].due_date == "2028-08-30"
    assert by_title["A007 Final Report — Final"].start_date == "2028-07-31"
    assert (
        by_title["A008 Patents - Reporting of Subject Inventions — Annual 1"].due_date
        == "2027-08-31"
    )
    assert (
        by_title["A008 Patents - Reporting of Subject Inventions — Annual 1"].start_date
        == "2026-08-31"
    )
    assert (
        by_title["A008 Patents - Reporting of Subject Inventions — Final"].due_date == "2028-08-30"
    )
    assert (
        by_title["A008 Patents - Reporting of Subject Inventions — Final"].start_date
        == "2027-08-31"
    )
    assert by_title["A009 One-Time Demonstration"].start_date == "2026-10-15"
    assert by_title["A009 One-Time Demonstration"].due_date == "2026-10-15"
    assert all("contract PoP 2026-08-31 to 2028-08-30" in row.notes for row in rows)
    assert all("Start inferred from" in row.notes for row in rows)


def test_cdrl_proposal_replaces_generic_templates(client: TestClient) -> None:
    admin = login(client)
    award = _award(
        client, admin, short_code="P12CDRL", type_code="FFP", template="FFP_INTERNAL", oh_pct=0
    )
    award_id = award["award_id"]
    created = client.post(
        f"/awards/{award_id}/documents",
        json={"kind_code": "contract", "title": "CDRL contract"},
        headers=auth_header(admin),
    )
    document_id = created.json()["document_id"]
    contract = b"""\
Section F - Deliveries or Performance
0001 Period of Performance
From
31 Aug 2026
To
30 Aug 2028
1. DATA ITEM NO.
A001
2. TITLE OF DATA ITEM
Program Management Plan
10. FREQUENCY
One time
12. DATE OF FIRST SUBMISSION
45 DAC
16. REMARKS
DAC: Days After Contract Award.
15.TOTAL
"""
    uploaded = client.post(
        f"/documents/{document_id}/file",
        files={"file": ("contract.txt", contract, "text/plain")},
        headers=auth_header(admin),
    )
    assert uploaded.status_code == 200, uploaded.text

    proposed = client.post(
        f"/awards/{award_id}/schedule/propose",
        json={"document_id": document_id},
        headers=auth_header(admin),
    )
    assert proposed.status_code == 200, proposed.text
    body = proposed.json()
    assert [row["title"] for row in body["items"]] == [
        "Contract period of performance",
        "A001 Program Management Plan",
    ]
    assert body["items"][0]["start_date"] == "2026-08-31"
    assert body["items"][0]["due_date"] == "2028-08-30"
    assert body["items"][1]["start_date"] == "2026-08-31"
    assert body["items"][1]["due_date"] == "2026-10-15"
    assert any("phase-template rows were replaced" in note for note in body["notes"])
    assert any("differs from the award record" in note for note in body["notes"])


def test_extract_from_docx_table(client: TestClient) -> None:
    admin = login(client)
    award = _award(
        client, admin, short_code="P12DOCX", type_code="FFP", template="FFP_INTERNAL", oh_pct=0
    )
    award_id = award["award_id"]
    created = client.post(
        f"/awards/{award_id}/documents",
        json={"kind_code": "contract", "title": "Word SOW"},
        headers=auth_header(admin),
    )
    doc_id = created.json()["document_id"]

    document = WordDocument()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Deliverable: Prototype review"
    table.cell(0, 1).text = "Due 10/15/2026"
    content = BytesIO()
    document.save(content)
    upload = client.post(
        f"/documents/{doc_id}/file",
        files={
            "file": (
                "sow.docx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=auth_header(admin),
    )
    assert upload.status_code == 200, upload.text

    proposed = client.post(
        f"/awards/{award_id}/schedule/propose",
        json={"document_id": doc_id},
        headers=auth_header(admin),
    )
    assert proposed.status_code == 200, proposed.text
    extracted = [row for row in proposed.json()["items"] if row["origin_code"] == "extract"]
    assert [row["due_date"] for row in extracted] == ["2026-10-15"]


def test_pdf_text_and_scanned_pdf_feedback(tmp_path, monkeypatch) -> None:
    path = tmp_path / "contract.pdf"
    path.write_bytes(b"%PDF fake test input")
    row = SimpleNamespace(stored_ext=".pdf")
    monkeypatch.setattr(contract_schedule, "stored_path", lambda _row: path)

    class TextPage:
        def extract_text(self) -> str:
            return "Deliverable A due 2026-05-01"

    monkeypatch.setattr(
        contract_schedule,
        "PdfReader",
        lambda _path: SimpleNamespace(pages=[TextPage()]),
    )
    text, note = contract_schedule._read_document_text(row)
    assert text == "Deliverable A due 2026-05-01"
    assert note is None

    class ScannedPage:
        def extract_text(self) -> None:
            return None

    monkeypatch.setattr(
        contract_schedule,
        "PdfReader",
        lambda _path: SimpleNamespace(pages=[ScannedPage()]),
    )
    text, note = contract_schedule._read_document_text(row)
    assert text is None
    assert "scanned" in note


def test_pasted_text_without_schedule_rows_returns_note(client: TestClient) -> None:
    admin = login(client)
    award = _award(
        client, admin, short_code="P12NOTE", type_code="FFP", template="FFP_INTERNAL", oh_pct=0
    )
    proposed = client.post(
        f"/awards/{award['award_id']}/schedule/propose",
        json={"text": "General contract narrative without a schedule date."},
        headers=auth_header(admin),
    )
    assert proposed.status_code == 200, proposed.text
    assert "no dated deliverable lines found in the pasted text" in proposed.json()["notes"]
